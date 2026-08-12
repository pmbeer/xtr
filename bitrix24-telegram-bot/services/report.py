from __future__ import annotations

from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from services.models import AnalyticsReport


class ReportGenerator:
    def __init__(self, output_dir: Path) -> None:
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate_excel(self, report: AnalyticsReport, *, filename: Optional[str] = None) -> Path:
        stamp = report.generated_at.strftime("%Y%m%d_%H%M%S")
        filename = filename or f"bitrix_report_{report.user_id}_{stamp}.xlsx"
        path = self.output_dir / filename

        wb = Workbook()
        ws_summary = wb.active
        ws_summary.title = "Сводка"

        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        section_font = Font(bold=True, size=12, color="1F4E79")

        ws_summary["A1"] = "Аналитика Bitrix24"
        ws_summary["A1"].font = Font(bold=True, size=16, color="1F4E79")
        ws_summary["A2"] = f"Сформировано: {report.generated_at:%d.%m.%Y %H:%M}"
        ws_summary["A3"] = f"Сотрудник: {report.user_name} (ID {report.user_id})"
        ws_summary["A4"] = (
            f"Период: {report.period.date_from:%d.%m.%Y} — {report.period.date_to:%d.%m.%Y}"
        )

        row = 6
        ws_summary[f"A{row}"] = "Показатель"
        ws_summary[f"B{row}"] = "Значение"
        for col in ("A", "B"):
            cell = ws_summary[f"{col}{row}"]
            cell.fill = header_fill
            cell.font = header_font

        t = report.tasks
        o = report.open_lines
        rows = [
            ("Задачи: всего", t.total),
            ("Задачи: закрыто", t.closed),
            ("Задачи: в работе", t.in_progress),
            ("Задачи: ждут выполнения", t.pending),
            ("Задачи: ждут контроля", t.awaiting_control),
            ("Задачи: отложены", t.deferred),
            ("Задачи: просрочены", t.overdue),
            (
                "Задачи: среднее время закрытия, ч",
                round(t.avg_close_hours, 2) if t.avg_close_hours is not None else "—",
            ),
            ("ОЛ: обработано сессий", o.total_sessions),
            ("ОЛ: закрыто", o.closed_sessions),
            ("ОЛ: в работе", o.active_sessions),
            (
                "ОЛ: среднее время первого ответа, сек",
                round(o.avg_answer_seconds, 1) if o.avg_answer_seconds is not None else "—",
            ),
            (
                "ОЛ: среднее время решения, сек",
                round(o.avg_close_seconds, 1) if o.avg_close_seconds is not None else "—",
            ),
            (
                "ОЛ: среднее время решения, ч",
                round(o.avg_close_hours, 2) if o.avg_close_hours is not None else "—",
            ),
        ]

        for metric, value in rows:
            row += 1
            ws_summary[f"A{row}"] = metric
            ws_summary[f"B{row}"] = value

        row += 2
        ws_summary[f"A{row}"] = "Текстовая сводка"
        ws_summary[f"A{row}"].font = section_font
        for line in report.summary_lines or report.build_summary():
            row += 1
            ws_summary[f"A{row}"] = line

        ws_summary.column_dimensions["A"].width = 48
        ws_summary.column_dimensions["B"].width = 24

        # Статусы
        ws_status = wb.create_sheet("Статусы задач")
        ws_status.append(["Статус", "Количество"])
        for cell in ws_status[1]:
            cell.fill = header_fill
            cell.font = header_font
        for status, count in sorted(t.by_status.items(), key=lambda x: (-x[1], x[0])):
            ws_status.append([status, count])
        ws_status.column_dimensions["A"].width = 28
        ws_status.column_dimensions["B"].width = 14

        # Стадии
        ws_stages = wb.create_sheet("Стадии")
        ws_stages.append(["Стадия", "Количество"])
        for cell in ws_stages[1]:
            cell.fill = header_fill
            cell.font = header_font
        for stage, count in sorted(t.by_stage.items(), key=lambda x: (-x[1], x[0])):
            ws_stages.append([stage, count])
        ws_stages.column_dimensions["A"].width = 32
        ws_stages.column_dimensions["B"].width = 14

        # Списки задач
        ws_lists = wb.create_sheet("Списки задач")
        ws_lists.append(["Тип", "Название"])
        for cell in ws_lists[1]:
            cell.fill = header_fill
            cell.font = header_font
        for title in t.closed_titles:
            ws_lists.append(["Закрыта", title])
        for title in t.in_progress_titles:
            ws_lists.append(["В работе", title])
        ws_lists.column_dimensions["A"].width = 14
        ws_lists.column_dimensions["B"].width = 70
        for col in ws_lists.columns:
            for cell in col:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Каналы ОЛ
        ws_ol = wb.create_sheet("Открытые линии")
        ws_ol.append(["Канал", "Сессий"])
        for cell in ws_ol[1]:
            cell.fill = header_fill
            cell.font = header_font
        if o.by_source:
            for source, count in sorted(o.by_source.items(), key=lambda x: (-x[1], x[0])):
                ws_ol.append([source, count])
        else:
            ws_ol.append(["Нет данных по каналам", 0])
        if o.notes:
            ws_ol.append([])
            ws_ol.append(["Примечание", o.notes])
        ws_ol.column_dimensions["A"].width = 28
        ws_ol.column_dimensions["B"].width = 60

        wb.save(path)
        return path

    def generate_txt(self, report: AnalyticsReport, *, filename: Optional[str] = None) -> Path:
        stamp = report.generated_at.strftime("%Y%m%d_%H%M%S")
        filename = filename or f"bitrix_report_{report.user_id}_{stamp}.txt"
        path = self.output_dir / filename
        lines = report.summary_lines or report.build_summary()
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        return path
