"""Генерация отчётов в Excel."""

from __future__ import annotations

import os
import tempfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from bitrix.analytics import UserReport


def _format_duration(seconds: float) -> str:
    if seconds < 60:
        return f"{seconds:.0f} сек"
    if seconds < 3600:
        return f"{seconds / 60:.1f} мин"
    return f"{seconds / 3600:.1f} ч"


def _format_period(report: UserReport) -> str:
    if report.period_from and report.period_to:
        return (
            f"{report.period_from.strftime('%d.%m.%Y')} — "
            f"{report.period_to.strftime('%d.%m.%Y')}"
        )
    return "Весь период"


def generate_excel_report(report: UserReport) -> str:
    """Создать Excel-файл с отчётом и вернуть путь к файлу."""
    wb = Workbook()

    # Сводка
    ws_summary = wb.active
    ws_summary.title = "Сводка"
    _write_summary(ws_summary, report)

    # Задачи
    ws_tasks = wb.create_sheet("Задачи")
    _write_tasks(ws_tasks, report)

    # Открытая линия
    ws_openlines = wb.create_sheet("Открытая линия")
    _write_openlines(ws_openlines, report)

    # Статусы задач
    ws_status = wb.create_sheet("Статусы задач")
    _write_status_breakdown(ws_status, report)

    tmp_dir = tempfile.gettempdir()
    filename = (
        f"bitrix_report_{report.user_id}_"
        f"{report.generated_at.strftime('%Y%m%d_%H%M%S')}.xlsx"
    )
    filepath = os.path.join(tmp_dir, filename)
    wb.save(filepath)
    return filepath


def _header_style(cell) -> None:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.alignment = Alignment(horizontal="center")


def _write_summary(ws, report: UserReport) -> None:
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 25

    rows = [
        ("Отчёт по активности Bitrix24", ""),
        ("Пользователь", report.user_name),
        ("ID пользователя", report.user_id),
        ("Период", _format_period(report)),
        ("Дата формирования", report.generated_at.strftime("%d.%m.%Y %H:%M")),
        ("", ""),
        ("ЗАДАЧИ", ""),
        ("Всего задач", report.tasks.total),
        ("Закрыто / завершено", report.tasks.closed),
        ("В работе", report.tasks.in_progress),
        ("", ""),
        ("ОТКРЫТАЯ ЛИНИЯ", ""),
        ("Всего обращений", report.open_lines.total_sessions),
        ("Закрытых обращений", report.open_lines.closed_sessions),
        ("Открытых обращений", report.open_lines.open_sessions),
        (
            "Среднее время решения",
            _format_duration(report.open_lines.avg_resolution_seconds)
            if report.open_lines.avg_resolution_seconds
            else "—",
        ),
        (
            "Среднее время первого ответа",
            _format_duration(report.open_lines.avg_first_response_seconds)
            if report.open_lines.avg_first_response_seconds
            else "—",
        ),
    ]

    for i, (label, value) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
        if label in ("ЗАДАЧИ", "ОТКРЫТАЯ ЛИНИЯ", "Отчёт по активности Bitrix24"):
            ws.cell(row=i, column=1).font = Font(bold=True, size=12)


def _write_tasks(ws, report: UserReport) -> None:
    headers = ["ID", "Название", "Статус", "Стадия", "Создана", "Закрыта", "Дедлайн"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        _header_style(cell)

    for row_idx, task in enumerate(report.tasks.tasks, start=2):
        status = int(task.get("status", task.get("STATUS", 0)))
        from bitrix.analytics import TASK_STATUS_NAMES

        ws.cell(row=row_idx, column=1, value=task.get("id", task.get("ID")))
        ws.cell(row=row_idx, column=2, value=task.get("title", task.get("TITLE")))
        ws.cell(row=row_idx, column=3, value=TASK_STATUS_NAMES.get(status, status))
        ws.cell(row=row_idx, column=4, value=task.get("stageId", task.get("STAGE_ID")))
        ws.cell(row=row_idx, column=5, value=task.get("createdDate", task.get("CREATED_DATE")))
        ws.cell(row=row_idx, column=6, value=task.get("closedDate", task.get("CLOSED_DATE")))
        ws.cell(row=row_idx, column=7, value=task.get("deadline", task.get("DEADLINE")))


def _write_openlines(ws, report: UserReport) -> None:
    headers = [
        "ID",
        "Статус",
        "Создана",
        "Закрыта",
        "Первый ответ",
        "Причина закрытия",
        "Чат ID",
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        _header_style(cell)

    for row_idx, session in enumerate(report.open_lines.sessions, start=2):
        ws.cell(row=row_idx, column=1, value=session.get("ID", session.get("id")))
        ws.cell(row=row_idx, column=2, value=session.get("STATUS", session.get("status")))
        ws.cell(
            row=row_idx,
            column=3,
            value=session.get("DATE_CREATE", session.get("dateCreate", session.get("CREATED"))),
        )
        ws.cell(
            row=row_idx,
            column=4,
            value=session.get("DATE_CLOSE", session.get("dateClose", session.get("END_TIME"))),
        )
        ws.cell(
            row=row_idx,
            column=5,
            value=session.get("DATE_OPERATOR", session.get("dateOperatorAnswer")),
        )
        ws.cell(
            row=row_idx,
            column=6,
            value=session.get("CLOSE_REASON", session.get("closeReason")),
        )
        ws.cell(row=row_idx, column=7, value=session.get("CHAT_ID", session.get("chatId")))


def _write_status_breakdown(ws, report: UserReport) -> None:
    ws.cell(row=1, column=1, value="Статус").font = Font(bold=True)
    ws.cell(row=1, column=2, value="Количество").font = Font(bold=True)

    row = 2
    for status_name, count in sorted(report.tasks.by_status.items()):
        ws.cell(row=row, column=1, value=status_name)
        ws.cell(row=row, column=2, value=count)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Стадия").font = Font(bold=True)
    ws.cell(row=row, column=2, value="Количество").font = Font(bold=True)
    row += 1

    for stage, count in sorted(report.tasks.by_stage.items()):
        ws.cell(row=row, column=1, value=stage)
        ws.cell(row=row, column=2, value=count)
        row += 1
