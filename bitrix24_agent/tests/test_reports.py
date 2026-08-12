from __future__ import annotations

import json
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from b24agent.analytics import period as periods
from b24agent.analytics.models import (
    ALL_SECTIONS,
    Report,
    ReportFormat,
    ReportRequest,
    ReportSection,
)
from b24agent.analytics.openlines_metrics import compute_openline_metrics
from b24agent.analytics.tasks_metrics import compute_task_metrics
from b24agent.bitrix.openlines import OpenLinesData, OpenLinesDataSource, Session
from b24agent.bitrix.tasks import Task, TaskStage
from b24agent.reports import build_report_file, build_summary
from b24agent.reports.builder import build_filename
from tests.fake_portal import session_payload, task_payload


@pytest.fixture
def report(tz, now) -> Report:
    period = periods.last_week(now)
    closed = [
        Task.from_api(
            task_payload(
                index,
                created="2026-08-03T10:00:00+03:00",
                closed="2026-08-03T14:00:00+03:00",
                deadline="2026-08-05T18:00:00+03:00",
                spent=3600,
            ),
            tz,
        )
        for index in (1, 2)
    ]
    open_tasks = [
        Task.from_api(task_payload(10, status=3, closed=None, stage_id=11), tz),
        Task.from_api(task_payload(11, status=2, closed=None), tz),
    ]
    tasks = compute_task_metrics(
        closed_tasks=closed,
        created_tasks=closed,
        open_tasks=open_tasks,
        period=period,
        stages={11: TaskStage(id=11, title="В работе", entity_id=0)},
        group_names={},
    )
    lines = compute_openline_metrics(
        OpenLinesData(
            source=OpenLinesDataSource.STATS_V2,
            sessions=[
                Session.from_v2(session_payload(index), tz) for index in (100, 101)
            ],
            line_names={3: "Поддержка"},
        ),
        period,
    )
    return Report(
        request=ReportRequest(
            period=period,
            sections=ALL_SECTIONS,
            report_format=ReportFormat.XLSX,
            raw_query="сколько задач я закрыл за прошлую неделю",
        ),
        period=period,
        generated_at=datetime(2026, 8, 12, 15, 30, tzinfo=tz),
        user_id=1,
        user_name="Иван Петров",
        portal_url="https://example.bitrix24.ru",
        tasks=tasks,
        openlines=lines,
        warnings=["Проверьте права вебхука"],
    )


def test_xlsx_has_expected_sheets_and_values(report):
    document = build_report_file(report)
    assert document.filename.endswith(".xlsx")
    assert document.content[:2] == b"PK"  # xlsx — это zip-архив

    workbook = load_workbook(BytesIO(document.content))
    assert workbook.sheetnames == [
        "Сводка",
        "Задачи по дням",
        "Стадии и статусы",
        "Закрытые задачи",
        "Задачи в работе",
        "Открытые линии",
        "Обращения",
    ]

    summary = workbook["Сводка"]
    labels = {row[0]: row[1] for row in summary.iter_rows(values_only=True) if row[0]}
    assert labels["Сотрудник"] == "Иван Петров"
    assert labels["Закрыто за период"] == 2
    assert labels["Сейчас не закрыто (всего)"] == 2
    assert labels["Среднее время закрытия задачи"] == "4 ч"
    assert labels["Всего обращений"] == 2

    closed_sheet = workbook["Закрытые задачи"]
    rows = list(closed_sheet.iter_rows(min_row=3, values_only=True))
    assert rows[0][0] == 1
    assert rows[0][5] == 4.0  # время закрытия в часах


def test_xlsx_without_details_skips_detail_sheets(report):
    report.request = ReportRequest(
        period=report.period,
        sections=ALL_SECTIONS,
        include_details=False,
    )
    workbook = load_workbook(BytesIO(build_report_file(report).content))
    assert "Закрытые задачи" not in workbook.sheetnames
    assert "Обращения" not in workbook.sheetnames
    assert "Сводка" in workbook.sheetnames


def test_xlsx_only_openlines_section(report):
    report.request = ReportRequest(
        period=report.period, sections=frozenset({ReportSection.OPENLINES})
    )
    report.tasks = None
    workbook = load_workbook(BytesIO(build_report_file(report).content))
    assert "Задачи по дням" not in workbook.sheetnames
    assert "Открытые линии" in workbook.sheetnames


def test_csv_is_excel_friendly(report):
    report.request = ReportRequest(period=report.period, report_format=ReportFormat.CSV)
    document = build_report_file(report)
    text = document.content.decode("utf-8")

    assert text.startswith("\ufeff")  # BOM, иначе Excel ломает кириллицу
    assert "Закрыто за период;2" in text
    assert "Задачи: закрыто по дням" in text
    assert "Открытые линии" in text
    assert document.filename.endswith(".csv")


def test_json_contains_meta_and_details(report):
    report.request = ReportRequest(period=report.period, report_format=ReportFormat.JSON)
    payload = json.loads(build_report_file(report).content)

    assert payload["meta"]["employee"] == {"id": 1, "name": "Иван Петров"}
    assert payload["meta"]["period"]["days"] == 7
    assert payload["tasks"]["closed_count"] == 2
    assert payload["openlines"]["total_sessions"] == 2
    assert len(payload["details"]["closed_tasks"]) == 2
    assert payload["details"]["open_tasks"][0]["stage"] == "В работе"


def test_summary_mentions_key_numbers(report):
    summary = build_summary(report)
    assert "Иван Петров" in summary
    assert "закрыто: <b>2</b>" in summary
    assert "среднее время решения" in summary
    assert "Проверьте права вебхука" in summary


def test_summary_escapes_html(report):
    report.user_name = "Иван <b>Петров</b>"
    assert "&lt;b&gt;" in build_summary(report)


def test_filename_is_transliterated(report):
    assert build_filename(report) == "bitrix24_ivan_petrov_20260803-20260809.xlsx"
