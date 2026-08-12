"""Отчёт в формате Excel: сводка, динамика и детализация по листам."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from b24agent.analytics.humanize import (
    date_ru,
    datetime_ru,
    duration,
    hours,
    percent,
    plural,
)
from b24agent.analytics.models import OpenLineMetrics, Report, ReportSection, TaskMetrics
from b24agent.bitrix.openlines import source_title
from b24agent.bitrix.tasks import status_name
from b24agent.reports.styles import (
    add_autofilter,
    set_widths,
    write_header,
    write_metric,
    write_note,
    write_row,
    write_section,
    write_title,
)

VOTE_LABELS = {"like": "Хорошо", "dislike": "Плохо", "none": "Нет оценки", "": "Нет оценки"}
STATUS_LABELS = {
    "new": "Новое",
    "answered": "В работе",
    "closed": "Закрыто",
    "spam": "Спам",
    "paused": "Приостановлено",
}
CLOSE_REASON_LABELS = {
    "operator": "Закрыл оператор",
    "auto": "Закрыто автоматически",
    "spam": "Спам",
    "client": "Закрыл клиент",
    "replyLimit": "Лимит ответов",
}


def render_xlsx(report: Report) -> bytes:
    workbook = Workbook()
    _summary_sheet(workbook.active, report)

    if report.tasks is not None:
        _tasks_dynamics_sheet(workbook.create_sheet("Задачи по дням"), report)
        if report.request.wants(ReportSection.STAGES):
            _stages_sheet(workbook.create_sheet("Стадии и статусы"), report.tasks)
        if report.request.include_details:
            _closed_tasks_sheet(workbook.create_sheet("Закрытые задачи"), report)
            _open_tasks_sheet(workbook.create_sheet("Задачи в работе"), report)

    if report.openlines is not None:
        _openlines_sheet(workbook.create_sheet("Открытые линии"), report.openlines)
        if report.request.include_details and report.openlines.sessions:
            _sessions_sheet(workbook.create_sheet("Обращения"), report.openlines)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


# ------------------------------------------------------------------- листы


def _summary_sheet(sheet: Worksheet, report: Report) -> None:
    sheet.title = "Сводка"
    set_widths(sheet, [46, 22, 58])

    row = write_title(sheet, 1, "Отчёт по активности в Битрикс24", width=3)
    row = write_metric(sheet, row, "Сотрудник", report.user_name)
    row = write_metric(sheet, row, "ID сотрудника", report.user_id)
    row = write_metric(sheet, row, "Период", report.period.human())
    row = write_metric(sheet, row, "Отчёт сформирован", datetime_ru(report.generated_at))
    if report.request.raw_query:
        row = write_metric(sheet, row, "Запрос", report.request.raw_query)
    row += 1

    if report.tasks is not None:
        row = _tasks_summary_block(sheet, row, report.tasks)
    if report.openlines is not None:
        row = _openlines_summary_block(sheet, row, report.openlines)

    if report.warnings:
        row = write_section(sheet, row, "Примечания", width=3)
        for warning in report.warnings:
            row = write_note(sheet, row, f"• {warning}", width=3)


def _tasks_summary_block(sheet: Worksheet, row: int, tasks: TaskMetrics) -> int:
    row = write_section(sheet, row, "Задачи", width=3)
    row = write_metric(sheet, row, "Закрыто за период", tasks.closed_count)
    row = write_metric(sheet, row, "Поставлено за период", tasks.created_count)
    row = write_metric(
        sheet, row, "Сейчас не закрыто (всего)", tasks.open_count, "срез на момент отчёта"
    )
    for status, count in tasks.by_status.items():
        row = write_metric(sheet, row, f"    из них «{status}»", count)
    row = write_metric(sheet, row, "Просрочено из незакрытых", tasks.overdue_open_count)
    row = write_metric(
        sheet,
        row,
        "Закрыто в срок / с просрочкой",
        f"{tasks.closed_on_time_count} / {tasks.closed_overdue_count}",
        f"без дедлайна: {tasks.without_deadline_count}",
    )
    row = write_metric(
        sheet,
        row,
        "Среднее время закрытия задачи",
        duration(tasks.avg_lead_time_seconds),
        "от постановки до завершения",
    )
    row = write_metric(
        sheet, row, "Медианное время закрытия", duration(tasks.median_lead_time_seconds)
    )
    row = write_metric(
        sheet,
        row,
        "90-й персентиль времени закрытия",
        duration(tasks.p90_lead_time_seconds),
        "9 из 10 задач закрываются быстрее",
    )
    row = write_metric(
        sheet, row, "Самая быстрая задача", duration(tasks.fastest_lead_time_seconds)
    )
    row = write_metric(
        sheet, row, "Самая долгая задача", duration(tasks.slowest_lead_time_seconds)
    )
    row = write_metric(
        sheet,
        row,
        "Списано времени по закрытым задачам",
        duration(tasks.total_time_spent_seconds or None),
        "по учёту рабочего времени в задачах",
    )
    return row + 1


def _openlines_summary_block(sheet: Worksheet, row: int, lines: OpenLineMetrics) -> int:
    row = write_section(sheet, row, "Открытые линии", width=3)
    row = write_metric(sheet, row, "Источник данных", lines.data_source.title)
    row = write_metric(sheet, row, "Всего обращений", lines.total_sessions)
    row = write_metric(sheet, row, "Закрыто обращений", lines.closed_sessions)
    row = write_metric(sheet, row, "Не закрыто обращений", lines.open_sessions)
    row = write_metric(
        sheet,
        row,
        "Среднее время до первого ответа",
        duration(lines.avg_first_answer_seconds),
    )
    row = write_metric(
        sheet,
        row,
        "Среднее время решения вопроса",
        duration(lines.avg_resolution_seconds),
        "от обращения клиента до закрытия",
    )
    row = write_metric(
        sheet,
        row,
        "Медианное время решения",
        duration(lines.median_resolution_seconds),
    )
    row = write_metric(
        sheet, row, "90-й персентиль времени решения", duration(lines.p90_resolution_seconds)
    )
    row = write_metric(
        sheet,
        row,
        "Оценки клиентов (хорошо / плохо)",
        f"{lines.likes} / {lines.dislikes}",
        f"доля положительных: {percent(lines.positive_rate)}",
    )
    if lines.kpi_first_answer_ok or lines.kpi_first_answer_fail:
        row = write_metric(
            sheet,
            row,
            "SLA первого ответа (уложились / нет)",
            f"{lines.kpi_first_answer_ok} / {lines.kpi_first_answer_fail}",
        )
    if lines.avg_messages:
        row = write_metric(
            sheet, row, "Среднее число сообщений в обращении", round(lines.avg_messages, 1)
        )
    if lines.by_source:
        row = write_section(sheet, row, "Обращения по каналам", width=3)
        for channel, count in lines.by_source.items():
            row = write_metric(sheet, row, f"    {channel}", count)
    if lines.by_line:
        row = write_section(sheet, row, "Обращения по линиям", width=3)
        for line_name, count in lines.by_line.items():
            row = write_metric(sheet, row, f"    {line_name}", count)
    return row + 1


def _tasks_dynamics_sheet(sheet: Worksheet, report: Report) -> None:
    tasks = report.tasks
    assert tasks is not None
    set_widths(sheet, [16, 22, 26])

    row = write_title(sheet, 1, "Динамика закрытия задач", width=3)
    row = write_header(sheet, row, ["Дата", "Закрыто задач", "Накопительно"])
    running = 0
    for day, count in tasks.closed_per_day.items():
        running += count
        row = write_row(sheet, row, [date_ru(day), count, running])
    row = write_row(sheet, row, ["Итого", tasks.closed_count, running])
    write_note(
        sheet,
        row + 1,
        "Задача попадает в день своего завершения (поле closedDate в Битрикс24).",
        width=3,
    )


def _stages_sheet(sheet: Worksheet, tasks: TaskMetrics) -> None:
    set_widths(sheet, [46, 16, 14])
    row = write_title(sheet, 1, "Незакрытые задачи: статусы, стадии, проекты", width=3)

    total = tasks.open_count or 1
    for caption, distribution in (
        ("По статусам", tasks.by_status),
        ("По стадиям канбана", tasks.by_stage),
        ("По проектам", tasks.by_group),
    ):
        row = write_section(sheet, row, caption, width=3)
        row = write_header(sheet, row, [caption, "Задач", "Доля"])
        for name, count in distribution.items():
            row = write_row(sheet, row, [name, count, percent(count / total)])
        if not distribution:
            row = write_row(sheet, row, ["Нет данных", 0, "—"])
        row += 1

    write_note(
        sheet,
        row,
        "Стадия — свойство задачи «здесь и сейчас»: REST Битрикс24 не отдаёт историю "
        "переходов между стадиями, поэтому распределение показано на момент отчёта.",
        width=3,
    )


def _closed_tasks_sheet(sheet: Worksheet, report: Report) -> None:
    tasks = report.tasks
    assert tasks is not None
    set_widths(sheet, [10, 60, 18, 18, 18, 16, 16, 14, 40])
    header = [
        "ID",
        "Название",
        "Поставлена",
        "Завершена",
        "Дедлайн",
        "Время закрытия, ч",
        "Списано, ч",
        "В срок",
        "Ссылка",
    ]
    row = write_title(sheet, 1, f"Закрытые задачи за период: {tasks.closed_count}", width=9)
    header_row = row
    row = write_header(sheet, row, header)
    for task in tasks.closed_tasks:
        row = write_row(
            sheet,
            row,
            [
                task.id,
                task.title,
                datetime_ru(task.created_at),
                datetime_ru(task.closed_at),
                datetime_ru(task.deadline),
                hours(task.lead_time_seconds),
                hours(task.time_spent),
                "нет" if task.is_overdue else "да",
                task.url(report.portal_url),
            ],
        )
    add_autofilter(sheet, header_row, len(header), row - 1)


def _open_tasks_sheet(sheet: Worksheet, report: Report) -> None:
    tasks = report.tasks
    assert tasks is not None
    set_widths(sheet, [10, 60, 18, 18, 18, 22, 14, 40])
    header = [
        "ID",
        "Название",
        "Статус",
        "Поставлена",
        "Дедлайн",
        "Стадия",
        "Просрочена",
        "Ссылка",
    ]
    row = write_title(sheet, 1, f"Задачи в работе: {tasks.open_count}", width=8)
    header_row = row
    row = write_header(sheet, row, header)

    for task in tasks.open_tasks:
        row = write_row(
            sheet,
            row,
            [
                task.id,
                task.title,
                status_name(task.status),
                datetime_ru(task.created_at),
                datetime_ru(task.deadline),
                _stage_label(tasks, task.stage_id),
                "да" if task.is_overdue else "нет",
                task.url(report.portal_url),
            ],
        )
    add_autofilter(sheet, header_row, len(header), row - 1)


def _openlines_sheet(sheet: Worksheet, lines: OpenLineMetrics) -> None:
    set_widths(sheet, [30, 18, 18, 30])
    row = write_title(sheet, 1, "Обращения открытых линий", width=4)

    if lines.is_approximate:
        row = write_note(
            sheet,
            row,
            "Цифры приблизительные: методы статистики открытых линий на портале "
            "недоступны, поэтому обращения посчитаны по делам CRM.",
            width=4,
        )
    row += 1

    row = write_section(sheet, row, "Обращения по дням", width=4)
    row = write_header(sheet, row, ["Дата", "Обращений", "Накопительно", ""])
    running = 0
    for day, count in lines.by_day.items():
        running += count
        row = write_row(sheet, row, [date_ru(day), count, running, ""])
    row += 1

    row = write_section(sheet, row, "Обращения по часам суток", width=4)
    row = write_header(sheet, row, ["Час", "Обращений", "", ""])
    for hour, count in enumerate(lines.by_hour):
        if count:
            row = write_row(sheet, row, [f"{hour:02d}:00", count, "", ""])
    row += 1

    row = write_section(sheet, row, "Обращения по каналам", width=4)
    row = write_header(sheet, row, ["Канал", "Обращений", "Доля", ""])
    total = lines.total_sessions or 1
    for channel, count in lines.by_source.items():
        row = write_row(sheet, row, [channel, count, percent(count / total), ""])


def _sessions_sheet(sheet: Worksheet, lines: OpenLineMetrics) -> None:
    set_widths(sheet, [12, 22, 26, 18, 18, 22, 22, 16, 14, 18])
    header = [
        "ID",
        "Канал",
        "Начало",
        "Первый ответ",
        "Закрыто",
        "До первого ответа, мин",
        "Время решения, мин",
        "Статус",
        "Оценка",
        "Сообщений",
    ]
    row = write_title(sheet, 1, f"Детализация обращений: {lines.total_sessions}", width=10)
    header_row = row
    row = write_header(sheet, row, header)

    for session in lines.sessions:
        row = write_row(
            sheet,
            row,
            [
                session.id,
                source_title(session.source),
                datetime_ru(session.created_at),
                datetime_ru(session.first_answer_at),
                datetime_ru(session.closed_at),
                _minutes(session.wait_answer_seconds),
                _minutes(session.resolution_seconds),
                STATUS_LABELS.get(session.status.lower(), session.status or "—"),
                VOTE_LABELS.get(session.vote, session.vote or "—"),
                session.message_count or "—",
            ],
        )
    add_autofilter(sheet, header_row, len(header), row - 1)


# ------------------------------------------------------------------ хелперы


def _stage_label(tasks: TaskMetrics, stage_id: int) -> str:
    if not stage_id:
        return "Без стадии"
    return tasks.stage_titles.get(stage_id, f"Стадия {stage_id}")


def _minutes(seconds: float | None) -> float | str:
    if seconds is None:
        return "—"
    return round(seconds / 60, 1)


def summary_counts(report: Report) -> str:
    """Короткая строка для подписи к файлу."""
    parts: list[str] = []
    if report.tasks is not None:
        closed = report.tasks.closed_count
        parts.append(
            f"закрыто {closed} {plural(closed, 'задача', 'задачи', 'задач')}"
        )
    if report.openlines is not None:
        total = report.openlines.total_sessions
        parts.append(
            f"обработано {total} {plural(total, 'обращение', 'обращения', 'обращений')}"
        )
    return ", ".join(parts)
