"""Оформление листов Excel: общие стили и хелперы записи."""

from __future__ import annotations

from collections.abc import Sequence
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ACCENT = "1F6FEB"
LIGHT = "EEF3FB"

TITLE_FONT = Font(bold=True, size=14, color="1B2733")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
SECTION_FONT = Font(bold=True, size=11, color=ACCENT)
SECTION_FILL = PatternFill("solid", fgColor=LIGHT)
MUTED_FONT = Font(italic=True, size=9, color="6B7785")
THIN_BORDER = Border(*(Side(style="thin", color="D6DEE8"),) * 4)


def write_title(sheet: Worksheet, row: int, text: str, width: int = 6) -> int:
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    sheet.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=max(width, 2)
    )
    return row + 1


def write_section(sheet: Worksheet, row: int, text: str, width: int = 6) -> int:
    for column in range(1, max(width, 2) + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = SECTION_FILL
        if column == 1:
            cell.value = text
            cell.font = SECTION_FONT
    return row + 1


def write_header(sheet: Worksheet, row: int, values: Sequence[str]) -> int:
    for index, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=index, value=value)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)
    return row + 1


def write_row(sheet: Worksheet, row: int, values: Sequence[Any], *, bordered: bool = True) -> int:
    for index, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=index, value=value)
        if bordered:
            cell.border = THIN_BORDER
        if isinstance(value, str) and len(value) > 60:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return row + 1


def write_metric(sheet: Worksheet, row: int, name: str, value: Any, note: str = "") -> int:
    sheet.cell(row=row, column=1, value=name).border = THIN_BORDER
    cell = sheet.cell(row=row, column=2, value=value)
    cell.border = THIN_BORDER
    cell.font = Font(bold=True)
    if note:
        hint = sheet.cell(row=row, column=3, value=note)
        hint.font = MUTED_FONT
    return row + 1


def write_note(sheet: Worksheet, row: int, text: str, width: int = 6) -> int:
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = MUTED_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(width, 2))
    return row + 1


def set_widths(sheet: Worksheet, widths: Sequence[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def add_autofilter(sheet: Worksheet, header_row: int, columns: int, last_row: int) -> None:
    if last_row <= header_row:
        return
    sheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(columns)}{last_row}"
    )
