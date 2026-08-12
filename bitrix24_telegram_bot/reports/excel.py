"""Формирование Excel-отчёта (xlsx) по собранной статистике."""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from analytics import OpenLinesStats, TaskStats
from analytics.periods import format_timedelta

HEADER_FILL = PatternFill("solid", fgColor="2FC6F6")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13)


def _fmt_dt(value: datetime | None) -> str:
    return value.strftime("%d.%m.%Y %H:%M") if value else ""


def _write_header(sheet: Worksheet, row: int, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autofit(sheet: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _fill_summary(
    sheet: Worksheet, tasks: TaskStats | None, openlines: OpenLinesStats | None
) -> None:
    period = (tasks or openlines).period  # хотя бы один раздел всегда есть
    sheet["A1"] = "Отчёт по активности в Bitrix24"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        f"Период: {period.label} "
        f"({period.date_from:%d.%m.%Y} — {period.date_to:%d.%m.%Y})"
    )
    sheet["A3"] = f"Сформирован: {datetime.now():%d.%m.%Y %H:%M}"

    row = 5
    _write_header(sheet, row, ["Показатель", "Значение"])
    rows: list[tuple[str, str]] = []

    if tasks is not None:
        rows += [
            ("Задач закрыто за период", str(tasks.closed_count)),
            ("Задач сейчас в работе (всего открытых)", str(tasks.open_count)),
            ("— из них со статусом «Выполняется»", str(tasks.in_progress_count)),
            ("— из них просрочено", str(tasks.overdue_count)),
            (
                "Среднее время выполнения задачи",
                format_timedelta(tasks.avg_resolution_seconds),
            ),
        ]
    if openlines is not None:
        rows += [
            ("Обращений (открытые линии) за период", str(openlines.total_count)),
            ("— из них обработано (завершено)", str(openlines.handled_count)),
            ("— из них ещё открыто", str(openlines.open_count)),
            (
                "Среднее время решения обращения",
                format_timedelta(openlines.avg_resolution_seconds),
            ),
        ]

    for label, value in rows:
        row += 1
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=value)

    if tasks is not None and tasks.open_by_stage:
        row += 2
        sheet.cell(row=row, column=1, value="Открытые задачи по стадиям").font = Font(
            bold=True
        )
        row += 1
        _write_header(sheet, row, ["Стадия / статус", "Задач"])
        for stage, count in tasks.open_by_stage.items():
            row += 1
            sheet.cell(row=row, column=1, value=stage)
            sheet.cell(row=row, column=2, value=count)

    _autofit(sheet, [45, 25])


def _fill_tasks_sheet(sheet: Worksheet, rows, include_closed: bool) -> None:
    headers = ["ID", "Задача", "Статус", "Стадия", "Создана", "Дедлайн"]
    if include_closed:
        headers += ["Закрыта", "Время выполнения"]
    _write_header(sheet, 1, headers)

    for index, task in enumerate(rows, start=2):
        values = [
            task.task_id,
            task.title,
            task.status_label,
            task.stage,
            _fmt_dt(task.created),
            _fmt_dt(task.deadline),
        ]
        if include_closed:
            values += [
                _fmt_dt(task.closed),
                format_timedelta(task.resolution_seconds),
            ]
        for col, value in enumerate(values, start=1):
            sheet.cell(row=index, column=col, value=value)

    _autofit(sheet, [10, 50, 18, 22, 18, 18, 18, 18][: len(headers)])
    sheet.freeze_panes = "A2"


def _fill_openlines_sheet(sheet: Worksheet, stats: OpenLinesStats) -> None:
    _write_header(
        sheet, 1, ["ID", "Обращение", "Создано", "Завершено", "Статус", "Время решения"]
    )
    for index, session in enumerate(stats.sessions, start=2):
        values = [
            session.session_id,
            session.subject,
            _fmt_dt(session.created),
            _fmt_dt(session.finished),
            "Обработано" if session.completed else "Открыто",
            format_timedelta(session.resolution_seconds),
        ]
        for col, value in enumerate(values, start=1):
            sheet.cell(row=index, column=col, value=value)

    _autofit(sheet, [10, 50, 18, 18, 14, 18])
    sheet.freeze_panes = "A2"


def build_excel_report(
    tasks: TaskStats | None, openlines: OpenLinesStats | None
) -> bytes:
    """Собрать xlsx-файл и вернуть его содержимое в байтах."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    _fill_summary(summary, tasks, openlines)

    if tasks is not None:
        closed_sheet = workbook.create_sheet("Закрытые задачи")
        _fill_tasks_sheet(closed_sheet, tasks.closed_tasks, include_closed=True)
        open_sheet = workbook.create_sheet("Задачи в работе")
        _fill_tasks_sheet(open_sheet, tasks.open_tasks, include_closed=False)

    if openlines is not None:
        openlines_sheet = workbook.create_sheet("Обращения")
        _fill_openlines_sheet(openlines_sheet, openlines)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
