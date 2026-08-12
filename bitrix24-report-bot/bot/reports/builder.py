"""Формирование Excel-файла с отчётом по активности в Битрикс24."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..bitrix.openlines import OpenLinesStats, format_duration
from ..bitrix.tasks import TaskStats
from ..utils.formatting import fmt_dt, fmt_hours, naive
from ..utils.period import Period

HEADER_FILL = PatternFill(start_color="FF2E5395", end_color="FF2E5395", fill_type="solid")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
SUBTITLE_FONT = Font(italic=True, color="FF666666")
METRIC_LABEL_FONT = Font(bold=True)


@dataclass
class ReportMeta:
    display_name: str
    period: Period
    generated_at: datetime


def _write_header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize_columns(ws: Worksheet, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _build_summary_sheet(
    wb: Workbook,
    meta: ReportMeta,
    task_stats: TaskStats,
    ol_stats: OpenLinesStats,
) -> None:
    ws = wb.active
    ws.title = "Сводка"

    ws["A1"] = f"Отчёт по активности в Битрикс24 — {meta.display_name}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Период: {meta.period.label}"
    ws["A2"].font = SUBTITLE_FONT
    ws["A3"] = f"Сформирован: {fmt_dt(naive(meta.generated_at))}"
    ws["A3"].font = SUBTITLE_FONT

    row = 5
    ws.cell(row=row, column=1, value="Задачи").font = Font(bold=True, size=12)
    row += 1
    metrics = [
        ("Закрыто задач за период", task_stats.closed_count),
        ("Создано задач за период", task_stats.created_count),
        ("В работе сейчас (всего)", task_stats.active_count),
        ("из них просрочено", task_stats.overdue_count),
        ("Среднее время выполнения задачи", fmt_hours(task_stats.avg_completion_hours)),
    ]
    for label, value in metrics:
        ws.cell(row=row, column=1, value=label).font = METRIC_LABEL_FONT
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Задачи в работе по стадиям").font = Font(bold=True, size=12)
    row += 1
    _write_header_row(ws, row, ["Стадия / статус", "Количество"])
    row += 1
    for stage, count in task_stats.by_stage.most_common():
        ws.cell(row=row, column=1, value=stage)
        ws.cell(row=row, column=2, value=count)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Открытые линии (обращения)").font = Font(bold=True, size=12)
    row += 1
    if ol_stats.error:
        ws.cell(row=row, column=1, value=ol_stats.error)
        row += 1
    else:
        ol_metrics = [
            ("Всего обращений за период", ol_stats.total),
            ("Обработано (закрыто)", ol_stats.completed),
            ("В работе / не закрыто", ol_stats.in_progress),
            ("Среднее время решения обращения", ol_stats.avg_resolution_human),
        ]
        for label, value in ol_metrics:
            ws.cell(row=row, column=1, value=label).font = METRIC_LABEL_FONT
            ws.cell(row=row, column=2, value=value)
            row += 1

    _autosize_columns(ws, [42, 22])


def _build_active_tasks_sheet(wb: Workbook, task_stats: TaskStats) -> None:
    ws = wb.create_sheet("Задачи в работе")
    headers = ["ID", "Название", "Статус", "Стадия", "Приоритет", "Создана", "Срок"]
    _write_header_row(ws, 1, headers)
    row = 2
    for task in task_stats.active_tasks:
        ws.cell(row=row, column=1, value=task.id)
        ws.cell(row=row, column=2, value=task.title)
        ws.cell(row=row, column=3, value=task.status_label)
        ws.cell(row=row, column=4, value=task.stage_label)
        ws.cell(row=row, column=5, value=task.priority_label)
        ws.cell(row=row, column=6, value=naive(task.created))
        ws.cell(row=row, column=7, value=naive(task.deadline))
        if task.is_overdue:
            for col in range(1, len(headers) + 1):
                ws.cell(row=row, column=col).fill = PatternFill(
                    start_color="FFFCE4E4", end_color="FFFCE4E4", fill_type="solid"
                )
        row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(row - 1, 1)}"
    _autosize_columns(ws, [8, 48, 18, 22, 12, 18, 18])


def _build_closed_tasks_sheet(wb: Workbook, task_stats: TaskStats) -> None:
    ws = wb.create_sheet("Закрытые задачи")
    headers = ["ID", "Название", "Стадия", "Создана", "Закрыта", "Время выполнения"]
    _write_header_row(ws, 1, headers)
    row = 2
    for task in task_stats.closed_tasks:
        duration = None
        if task.created and task.closed:
            duration = fmt_hours((task.closed - task.created).total_seconds() / 3600.0)
        ws.cell(row=row, column=1, value=task.id)
        ws.cell(row=row, column=2, value=task.title)
        ws.cell(row=row, column=3, value=task.stage_label)
        ws.cell(row=row, column=4, value=naive(task.created))
        ws.cell(row=row, column=5, value=naive(task.closed))
        ws.cell(row=row, column=6, value=duration or "—")
        row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(row - 1, 1)}"
    _autosize_columns(ws, [8, 48, 22, 18, 18, 18])


def _build_openlines_sheet(wb: Workbook, ol_stats: OpenLinesStats) -> None:
    ws = wb.create_sheet("Открытые линии")
    if ol_stats.error:
        ws["A1"] = ol_stats.error
        _autosize_columns(ws, [100])
        return

    headers = ["ID", "Тема обращения", "Создано", "Завершено", "Статус", "Время решения"]
    _write_header_row(ws, 1, headers)
    row = 2
    for session in ol_stats.sessions:
        duration = session.duration_seconds
        ws.cell(row=row, column=1, value=session.id)
        ws.cell(row=row, column=2, value=session.subject)
        ws.cell(row=row, column=3, value=naive(session.created))
        ws.cell(row=row, column=4, value=naive(session.end_time))
        ws.cell(row=row, column=5, value="Закрыто" if session.completed else "В работе")
        ws.cell(row=row, column=6, value=format_duration(duration) if duration is not None else "—")
        row += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(row - 1, 1)}"
    _autosize_columns(ws, [10, 48, 18, 18, 14, 16])


def build_report_workbook(
    meta: ReportMeta,
    task_stats: TaskStats,
    ol_stats: OpenLinesStats,
) -> Workbook:
    wb = Workbook()
    _build_summary_sheet(wb, meta, task_stats, ol_stats)
    _build_active_tasks_sheet(wb, task_stats)
    _build_closed_tasks_sheet(wb, task_stats)
    _build_openlines_sheet(wb, ol_stats)
    return wb


def build_report_file(
    output_path: str | Path,
    meta: ReportMeta,
    task_stats: TaskStats,
    ol_stats: OpenLinesStats,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = build_report_workbook(meta, task_stats, ol_stats)
    wb.save(output_path)
    return output_path
