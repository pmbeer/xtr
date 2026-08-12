from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from bitrix24_agent.analytics import (
    AnalyticsResult,
    OpenLinesAnalytics,
    TaskAnalytics,
)
from bitrix24_agent.periods import ReportPeriod
from bitrix24_agent.report import build_xlsx, format_duration


def test_builds_safe_xlsx_report() -> None:
    timezone = ZoneInfo("Europe/Moscow")
    result = AnalyticsResult(
        user_id=42,
        period=ReportPeriod(
            datetime(2026, 8, 1, tzinfo=timezone),
            datetime(2026, 8, 12, tzinfo=timezone),
            "01.08.2026–12.08.2026",
        ),
        tasks=TaskAnalytics(
            completed=[{"id": "1", "title": "=HYPERLINK(\"bad\")", "status": "5"}],
        ),
        open_lines=OpenLinesAnalytics(
            processed=4,
            closed=3,
            average_first_answer_seconds=65,
            average_resolution_seconds=3661,
        ),
    )

    workbook = load_workbook(BytesIO(build_xlsx(result)))
    assert workbook.sheetnames == [
        "Сводка",
        "Статусы и стадии",
        "Активные задачи",
        "Завершённые задачи",
    ]
    assert workbook["Завершённые задачи"]["B2"].value.startswith("'=")
    assert workbook["Сводка"]["B3"].value == 42


def test_duration_format() -> None:
    assert format_duration(3661) == "01:01:01"
    assert format_duration(None) == "Нет данных"
