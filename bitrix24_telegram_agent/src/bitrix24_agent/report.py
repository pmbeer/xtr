from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .analytics import STATUS_NAMES, AnalyticsResult, get_field

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def format_duration(seconds: float | None) -> str:
    if seconds is None:
        return "Нет данных"
    total = max(0, round(seconds))
    hours, remainder = divmod(total, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _header(sheet: Worksheet, row: int = 1) -> None:
    for cell in sheet[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autosize(sheet: Worksheet) -> None:
    for column in sheet.columns:
        letter = column[0].column_letter
        width = max((len(str(cell.value or "")) for cell in column), default=0)
        sheet.column_dimensions[letter].width = min(max(width + 2, 10), 55)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _task_rows(tasks: list[dict[str, Any]]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for task in tasks:
        status = int(get_field(task, "STATUS", 0) or 0)
        rows.append(
            [
                get_field(task, "ID"),
                safe_cell(get_field(task, "TITLE", "")),
                STATUS_NAMES.get(status, "Неизвестно"),
                get_field(task, "STAGE_ID", ""),
                get_field(task, "DEADLINE", ""),
                get_field(task, "CREATED_DATE", ""),
                get_field(task, "CLOSED_DATE", ""),
            ]
        )
    return rows


def build_xlsx(result: AnalyticsResult) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    summary.append(["Показатель", "Значение"])
    summary_rows = [
        ("Период", result.period.label),
        ("ID сотрудника Bitrix24", result.user_id),
        ("Завершено задач как исполнителем", len(result.tasks.completed)),
        ("Активных задач", len(result.tasks.active)),
        ("Просроченных активных задач", result.tasks.overdue),
        ("Обработано обращений", result.open_lines.processed),
        ("Закрыто обращений", result.open_lines.closed),
        ("Активных обращений", result.open_lines.active),
        (
            "Среднее время первого ответа",
            format_duration(result.open_lines.average_first_answer_seconds),
        ),
        (
            "Среднее время решения",
            format_duration(result.open_lines.average_resolution_seconds),
        ),
    ]
    for label, value in summary_rows:
        summary.append([label, "Нет данных" if value is None else value])
    if result.open_lines.error:
        summary.append(["Открытые линии", f"Недоступно: {result.open_lines.error}"])
    _header(summary)
    _autosize(summary)

    breakdown = workbook.create_sheet("Статусы и стадии")
    breakdown.append(["Тип", "Статус / стадия", "Количество"])
    for name, count in result.tasks.statuses.most_common():
        breakdown.append(["Статус", name, count])
    for name, count in result.tasks.stages.most_common():
        breakdown.append(["Стадия", name, count])
    _header(breakdown)
    _autosize(breakdown)

    headers = ["ID", "Название", "Статус", "ID стадии", "Дедлайн", "Создана", "Закрыта"]
    for title, tasks in (
        ("Активные задачи", result.tasks.active),
        ("Завершённые задачи", result.tasks.completed),
    ):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in _task_rows(tasks):
            sheet.append(row)
        _header(sheet)
        _autosize(sheet)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
